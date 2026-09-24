"""Create a sample DOCX with headings, text, and a table for testing."""

from __future__ import annotations

from pathlib import Path

from docx import Document

OUT = Path(__file__).resolve().parent.parent / "sample_document.docx"


def main() -> None:
    doc = Document()
    doc.add_heading("HR Policies Manual", 0)
    doc.add_heading("Leave Policy", level=1)
    doc.add_paragraph(
        "Employees are entitled to casual leave, annual leave, and sick leave "
        "as per company policy. Casual leave may be taken for personal matters."
    )
    doc.add_heading("Leave Entitlements", level=2)
    table = doc.add_table(rows=4, cols=3)
    headers = ["Leave Type", "Days Per Year", "Notes"]
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    rows = [
        ("Casual Leave", "12", "For personal matters"),
        ("Annual Leave", "21", "Accrued monthly"),
        ("Sick Leave", "10", "Medical certificate required"),
    ]
    for r_idx, row in enumerate(rows, start=1):
        for c_idx, val in enumerate(row):
            table.rows[r_idx].cells[c_idx].text = val

    doc.add_heading("Remote Work", level=1)
    doc.add_paragraph(
        "Remote work is permitted up to three days per week with manager approval."
    )

    doc.save(OUT)
    print(f"Created {OUT}")


if __name__ == "__main__":
    main()
